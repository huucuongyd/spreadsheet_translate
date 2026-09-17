"""
    Translates a Word document (.docx) or a folder of .docx files.

    >> ./docx_translate.py -n file.docx -c "ja" -t "vi"
    >> ./docx_translate.py -n docx/ -c "ja" -t "vi"

    Args:  -n   .docx file or folder containing .docx files
           -c   current language of the document (e.g. ja, en)
           -t   target language of the document

    from docx_translate import DocxDocument, translate_path

    document = DocxDocument("file.docx", current="ja", target="vi")
    document.translate()
    document.save()

    translate_path("docx/", current="ja", target="vi")

    IMPORTANT: Uses the free "translate" package (MyMemory). Large files may hit
    daily free-translation limits.
"""

import logging
import argparse
from pathlib import Path
from docx import Document
from translate import Translator


LEVEL = logging.INFO
FMT = '[%(levelname)s] %(asctime)s - %(message)s'
logging.basicConfig(level=LEVEL, format=FMT)

parser = argparse.ArgumentParser()
parser.add_argument("-n", "--name", type=str, required=True,
                    help="A .docx file or a folder containing .docx files")
parser.add_argument("-c", "--current", type=str, required=True, help="Current language of the document")
parser.add_argument("-t", "--to", type=str, required=True, help="Target language of the document")


class DocxDocument:
    """Opens, translates and saves a Word document (.docx)"""

    def __init__(self, filename_: str, current: str, target: str, translator: Translator | None = None):
        self.translator = translator or Translator(from_lang=current, to_lang=target)
        self.document = None
        self.filename = str(filename_)

    def translate_text(self, text: str) -> str:
        """Translates non-empty text; leaves blanks unchanged."""
        if text is None or not str(text).strip():
            return text
        return self.translator.translate(text)

    def translate_paragraph(self, paragraph) -> None:
        """Translates a paragraph while keeping the first run's formatting."""
        original = paragraph.text
        if not original.strip():
            return

        translated = self.translate_text(original)
        if not paragraph.runs:
            paragraph.add_run(translated)
            return

        paragraph.runs[0].text = translated
        for run in paragraph.runs[1:]:
            run.text = ""

    def translate_table(self, table) -> None:
        """Translates every cell in a table."""
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    self.translate_paragraph(paragraph)

    def translate_container(self, container) -> None:
        """Translates paragraphs and tables inside a document part."""
        for paragraph in container.paragraphs:
            self.translate_paragraph(paragraph)
        for table in container.tables:
            self.translate_table(table)

    def translate(self):
        """Loads the document and translates body, headers and footers."""
        logging.info('Loading "%s".', self.filename)
        try:
            self.document = Document(self.filename)
        except Exception:
            logging.error("Error occurred. Couldn't load the document: %s", self.filename)
            return None

        logging.info('Translation started for "%s"...', self.filename)
        self.translate_container(self.document)

        for section in self.document.sections:
            self.translate_container(section.header)
            self.translate_container(section.footer)

        logging.info('Document translation completed: "%s"', self.filename)
        return self.document

    def save(self, output_path: str | Path | None = None):
        """Saves the document into a new .docx file next to the original."""
        if output_path is None:
            path = Path(self.filename)
            output_path = path.with_name(f'Translated - {path.name}')
        else:
            output_path = Path(output_path)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        logging.info('Saving the document into "%s"', output_path)
        self.document.save(str(output_path))
        logging.info('Done: "%s"', output_path)
        return output_path


def collect_docx_files(path: Path) -> list[Path]:
    """Collect .docx files from a file path or a folder (recursive)."""
    if path.is_file():
        if path.suffix.lower() != '.docx':
            raise ValueError(f'Not a .docx file: {path}')
        return [path]

    if not path.is_dir():
        raise FileNotFoundError(f'Path not found: {path}')

    files = []
    for file_path in sorted(path.rglob('*.docx')):
        name = file_path.name
        if name.startswith('~$'):
            continue
        if name.startswith('Translated - '):
            continue
        files.append(file_path)
    return files


def translate_path(name: str, current: str, target: str) -> list[Path]:
    """Translate one .docx file or every .docx file under a folder."""
    path = Path(name)
    files = collect_docx_files(path)
    if not files:
        logging.warning('No .docx files found in "%s"', path)
        return []

    logging.info('Found %s .docx file(s) to translate.', len(files))
    translator = Translator(from_lang=current, to_lang=target)
    saved = []

    for file_path in files:
        document = DocxDocument(str(file_path), current=current, target=target, translator=translator)
        if document.translate() is None:
            continue
        saved.append(document.save())

    logging.info('Finished. Translated %s/%s file(s).', len(saved), len(files))
    return saved


if __name__ == "__main__":
    args = parser.parse_args()
    translate_path(args.name, current=args.current, target=args.to)
